"""Evidence-oriented document parsing and checks. No authenticity guesses."""
from __future__ import annotations

import io
import hashlib
import os
import re
import shutil
import subprocess
import json
from collections import Counter
from datetime import datetime
from pathlib import Path

MAX_PAGES = 20


def finding(name, status, message, evidence=None, severity="info"):
    return {"name": name, "status": status, "message": message,
            "evidence": evidence or [], "severity": severity}


def _metadata_analysis(data: bytes, file_type: str, embedded=None):
    """Inventory available embedded metadata; never infer that it is genuine."""
    embedded = embedded or {}
    unavailable = "Not present in embedded metadata"
    values = {
        "file_type": file_type,
        "creation_date": unavailable,
        "modified_date": unavailable,
        "author_creator": unavailable,
        "software_used": unavailable,
        "gps_location": "Not applicable / not present",
        "file_sha256": hashlib.sha256(data).hexdigest(),
        "hash_meaning": "Any byte change changes this fingerprint; it does not identify or authenticate the issuer",
        "original_metadata": "Cannot determine without a trusted original copy for comparison",
        "timestamp_limit": "Original computer file timestamps are not supplied to this service",
    }
    normalized = {str(k).lower().replace("_", ""): str(v).strip() for k, v in embedded.items() if v}
    def pick(*names):
        for name in names:
            for key, value in normalized.items():
                if name in key and value:
                    return value[:200]
        return unavailable
    values["creation_date"] = pick("datetimeoriginal", "createdate", "creationdate", "created")
    values["modified_date"] = pick("datetimemodified", "modifydate", "moddate", "modified", "datetime")
    values["author_creator"] = pick("author", "artist", "creator", "owner")
    values["software_used"] = pick("software", "creatortool", "creator_tool", "producer")
    gps = pick("gpsinfo", "gpslatitude", "gpslongitude", "gps")
    if gps != unavailable:
        values["gps_location"] = "Location metadata present (coordinates hidden)"
    return values


def _exiftool_metadata(data: bytes):
    """Read metadata with ExifTool when the optional local executable is installed."""
    configured = os.getenv("EXIFTOOL_CMD")
    executable = configured if configured and Path(configured).is_file() else shutil.which("exiftool")
    if not executable and os.name == "nt":
        candidate = Path(r"C:\Tools\ExifTool\exiftool.exe")
        if candidate.is_file():
            executable = str(candidate)
    if not executable:
        return {}, None
    try:
        result = subprocess.run([executable, "-json", "-G1", "-s", "-"], input=data,
                                capture_output=True, timeout=8, check=False)
        if result.returncode != 0:
            return {}, "ExifTool could not parse this image; built-in metadata reader was used."
        payload = json.loads(result.stdout.decode("utf-8", "replace"))
        if not isinstance(payload, list) or not payload or not isinstance(payload[0], dict):
            return {}, "ExifTool returned no metadata; built-in metadata reader was used."
        return {str(k): str(v)[:500] for k, v in payload[0].items()
                if k != "SourceFile" and not str(k).startswith("File:") and v is not None}, None
    except (OSError, subprocess.TimeoutExpired, ValueError):
        return {}, "ExifTool was unavailable for this image; built-in metadata reader was used."


def _image_metadata(data: bytes, image, extension: str):
    from PIL.ExifTags import TAGS
    embedded = {}
    for tag_id, value in image.getexif().items():
        tag = TAGS.get(tag_id, str(tag_id))
        if tag == "GPSInfo":
            embedded["GPSInfo"] = "present"
        else:
            embedded[tag] = str(value)[:200]
    # XMP commonly carries creator/software dates in addition to EXIF. Parse
    # bounded XML tags and attributes; never display the raw packet in reports.
    xmp = image.info.get("XML:com.adobe.xmp") or image.info.get("xmp")
    if xmp:
        if isinstance(xmp, bytes):
            xmp = xmp.decode("utf-8", "replace")
        embedded["XMP packet"] = "present"
        if len(xmp) <= 256_000 and "<!DOCTYPE" not in xmp.upper() and "<!ENTITY" not in xmp.upper():
            try:
                import xml.etree.ElementTree as ET
                root = ET.fromstring(xmp)
                for node in root.iter():
                    tag = node.tag.rsplit("}", 1)[-1].split(":")[-1]
                    if node.text and node.text.strip() and len(node) == 0:
                        embedded[f"XMP:{tag}"] = node.text.strip()[:300]
                    for attr, value in node.attrib.items():
                        attr = attr.rsplit("}", 1)[-1].split(":")[-1]
                        embedded[f"XMP:{attr}"] = str(value)[:300]
            except (ET.ParseError, ValueError):
                pass
        if "GPSLatitude" in xmp or "GPSLongitude" in xmp:
            embedded["GPS"] = "present"
    # Include standard image.info keys but keep the report bounded and avoid raw XMP dumps.
    for key, value in image.info.items():
        if key not in {"icc_profile", "exif", "XML:com.adobe.xmp", "xmp"} and isinstance(value, (str, int, float)):
            embedded[str(key)] = str(value)[:200]
    richer, tool_note = _exiftool_metadata(data)
    # Keep all useful tool tags for the detailed inventory, but avoid exposing
    # GPS coordinates in the browser/report.
    embedded.update(richer)
    result = _metadata_analysis(data, "JPEG image" if extension in {".jpg", ".jpeg"} else "PNG image", embedded)
    result["metadata_present"] = bool(embedded)
    result["metadata_keys"] = sorted(embedded.keys())
    details = []
    for key, value in sorted(embedded.items(), key=lambda item: item[0].lower()):
        if key.lower().find("gps") >= 0 or key.lower().find("latitude") >= 0 or key.lower().find("longitude") >= 0:
            value = "Location metadata present (coordinates hidden)"
        if key == "XMP packet":
            value = "Present; individual XMP fields are listed where readable"
        details.append({"field": key, "value": str(value)[:300]})
    result["metadata_details"] = details[:100]
    result["metadata_reader"] = "ExifTool + built-in reader" if richer else "Pillow built-in reader"
    result["metadata_reader_note"] = tool_note or ("Install ExifTool for broader EXIF, IPTC and XMP coverage." if not richer else "")
    return result


def _ocr_image_tesseract(image):
    try:
        import pytesseract
        from PIL import Image, ImageOps
        if os.getenv("TESSERACT_CMD"):
            pytesseract.pytesseract.tesseract_cmd = os.environ["TESSERACT_CMD"]
        elif not shutil.which("tesseract"):
            for candidate in (r"C:\Program Files\Tesseract-OCR\tesseract.exe",
                              r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"):
                if Path(candidate).is_file():
                    pytesseract.pytesseract.tesseract_cmd = candidate
                    break
        from collections import defaultdict
        from pytesseract import Output
        source = Image.open(io.BytesIO(image))
        source.thumbnail((3000, 3000), Image.Resampling.LANCZOS)
        raw = pytesseract.image_to_data(source, output_type=Output.DICT)
        grouped = defaultdict(list)
        heights, confidences, words = [], [], []
        for i, token in enumerate(raw["text"]):
            token = token.strip()
            if not token:
                continue
            try:
                confidence = float(raw["conf"][i])
            except (TypeError, ValueError):
                confidence = -1
            if confidence >= 0:
                confidences.append(confidence)
            if confidence < 15:
                continue
            key = (raw["block_num"][i], raw["par_num"][i], raw["line_num"][i])
            grouped[key].append(token)
            heights.append(int(raw["height"][i]))
            words.append({"text": token, "left": int(raw["left"][i]), "top": int(raw["top"][i]),
                          "width": int(raw["width"][i]), "height": int(raw["height"][i])})
        text = "\n".join(" ".join(tokens) for tokens in grouped.values())
        # Handwritten amount columns are often missed by prose OCR. Crop and
        # enlarge the numeric area, then try enhanced and thresholded scans.
        numeric_lines = []
        try:
            x0, x1 = int(source.width * .73), int(source.width * .97)
            y0, y1 = int(source.height * .27), int(source.height * .91)
            crop = ImageOps.autocontrast(source.convert("L").crop((x0, y0, x1, y1)))
            scale = max(2, min(4, 1500 // max(crop.width, crop.height)))
            enlarged = crop.resize((crop.width * scale, crop.height * scale))
            thresholded = enlarged.point(lambda p: 255 if p > 165 else 0)
            candidates = []
            for prepared in (enlarged, thresholded):
                for psm in (6, 11):
                    numeric = pytesseract.image_to_data(prepared, output_type=Output.DICT,
                        config=f"--psm {psm} -c tessedit_char_whitelist=0123456789,.")
                    tokens = []
                    for i, raw_token in enumerate(numeric["text"]):
                        token = raw_token.strip()
                        if not token or not re.fullmatch(r"[\d,.]+", token):
                            continue
                        try: conf = float(numeric["conf"][i])
                        except (TypeError, ValueError): conf = -1
                        tokens.append({"text": token, "left": x0 + int(numeric["left"][i]) / scale,
                            "top": y0 + int(numeric["top"][i]) / scale,
                            "confidence": conf, "height": max(1, int(numeric["height"][i]) / scale)})
                    # OCR can put rupees and paise in different text lines.
                    # Pair numeric tokens by their physical row position.
                    rows = []
                    for token in sorted(tokens, key=lambda t: (t["top"] + t["height"] / 2, t["left"])):
                        center = token["top"] + token["height"] / 2
                        match = next((row for row in reversed(rows)
                                      if abs(row["center"] - center) <= max(12, token["height"])), None)
                        if match is None:
                            rows.append({"center": center, "entries": [token]})
                        else:
                            match["entries"].append(token)
                            match["center"] = sum(x["top"] + x["height"] / 2 for x in match["entries"]) / len(match["entries"])
                    candidate = []
                    for row in rows:
                        entries = sorted(row["entries"], key=lambda t: t["left"])
                        candidate.append({"top": round(row["center"]),
                            "tokens": [x["text"] for x in entries],
                            "confidence": round(sum(x["confidence"] for x in entries) / len(entries), 1)})
                    candidates.append(candidate)
            numeric_lines = max(candidates, key=lambda rows: (sum(len(r["tokens"]) >= 2 for r in rows), len(rows)), default=[])
        except Exception:
            pass
        metrics = {"average_confidence": round(sum(confidences) / len(confidences), 1) if confidences else None,
                   "word_count": len(words), "width": source.width, "height": source.height,
                   "words": words, "line_count": len(grouped), "numeric_lines": numeric_lines}
        return text, metrics, None
    except ImportError as exc:
        missing = "Tesseract connector is not installed" if "pytesseract" in str(exc).lower() else "Pillow image support is not installed"
        return "", None, missing
    except Exception as exc:
        if type(exc).__name__ == "TesseractNotFoundError":
            return "", None, "Tesseract OCR executable is unavailable. Install it and set TESSERACT_CMD to its full path."
        return "", None, f"OCR could not complete ({type(exc).__name__}). Check the local OCR installation and image." 


def _analysis_raster_bytes(data: bytes, extension: str) -> bytes:
    """Create a bounded-size working raster while keeping metadata/hash on original bytes."""
    from PIL import Image, ImageOps
    source = Image.open(io.BytesIO(data))
    width, height = source.size
    pixel_count = width * height
    max_pixels = 60_000_000 if extension in {".jpg", ".jpeg"} else 30_000_000
    if pixel_count > max_pixels:
        raise ValueError(f"Image resolution ({width}×{height}) is too large for this demo service. Resize it and try again.")
    if extension in {".jpg", ".jpeg"}:
        # JPEG decoders can downsample during decode, preventing large phone
        # photos from expanding to hundreds of megabytes on a small free host.
        source.draft("RGB", (2600, 2600))
    source = ImageOps.exif_transpose(source).convert("RGB")
    source.thumbnail((2600, 2600), Image.Resampling.LANCZOS)
    output = io.BytesIO()
    if extension == ".png":
        source.save(output, format="PNG")
    else:
        source.save(output, format="JPEG", quality=90, optimize=True)
    return output.getvalue()


_PADDLE_ENGINE = None


def _ocr_image_paddle(image):
    """Optional PaddleOCR v3 local CPU/GPU engine; model downloads are managed by PaddleOCR."""
    global _PADDLE_ENGINE
    try:
        import numpy as np
        from PIL import Image
        from paddleocr import PaddleOCR
        if _PADDLE_ENGINE is None:
            _PADDLE_ENGINE = PaddleOCR(lang=os.getenv("PADDLEOCR_LANG", "en"),
                use_doc_orientation_classify=False, use_doc_unwarping=False,
                use_textline_orientation=True, engine="paddle")
        source = Image.open(io.BytesIO(image)).convert("RGB")
        source.thumbnail((3000, 3000), Image.Resampling.LANCZOS)
        output = _PADDLE_ENGINE.predict(input=np.asarray(source))
        tokens = []
        confidences = []
        for item in output or []:
            payload = getattr(item, "json", item)
            if callable(payload):
                payload = payload()
            if isinstance(payload, str):
                payload = json.loads(payload)
            result = payload.get("res", payload) if isinstance(payload, dict) else {}
            texts = result.get("rec_texts", []) or []
            scores = result.get("rec_scores", []) or []
            boxes = result.get("rec_boxes", []) or []
            for i, token in enumerate(texts):
                token = str(token).strip()
                if not token:
                    continue
                try:
                    score = float(scores[i])
                except (IndexError, TypeError, ValueError):
                    score = 0.0
                try:
                    x1, y1, x2, y2 = map(float, boxes[i])
                except (IndexError, TypeError, ValueError):
                    x1 = y1 = x2 = y2 = 0.0
                tokens.append({"text": token, "left": max(0, int(x1)), "top": max(0, int(y1)),
                    "width": max(1, int(x2 - x1)), "height": max(1, int(y2 - y1)), "score": score})
                confidences.append(score)
        tokens.sort(key=lambda w: (w["top"] // max(12, w["height"] // 2), w["left"]))
        lines = []
        for word in tokens:
            center = word["top"] + word["height"] / 2
            row = next((r for r in reversed(lines) if abs(r["center"] - center) <= max(12, word["height"] * .65)), None)
            if row is None:
                lines.append({"center": center, "words": [word]})
            else:
                row["words"].append(word)
                row["center"] = sum(w["top"] + w["height"] / 2 for w in row["words"]) / len(row["words"])
        lines.sort(key=lambda r: r["center"])
        for row in lines:
            row["words"].sort(key=lambda w: w["left"])
        text = "\n".join(" ".join(w["text"] for w in row["words"]) for row in lines)
        words = [{k: v for k, v in word.items() if k != "score"} for word in tokens]
        metrics = {"average_confidence": round(100 * sum(confidences) / len(confidences), 1) if confidences else None,
            "word_count": len(words), "width": source.width, "height": source.height, "words": words,
            "line_count": len(lines), "numeric_lines": [], "engine": "PaddleOCR"}
        return text, metrics, None
    except ImportError:
        return "", None, "PaddleOCR is not installed. Install the optional vision requirements or configure Tesseract."
    except Exception as exc:
        return "", None, f"PaddleOCR could not complete ({type(exc).__name__}); check its local model/runtime setup."


def _ocr_image(image):
    engine = os.getenv("OCR_ENGINE", "auto").strip().lower()
    if engine == "paddle":
        return _ocr_image_paddle(image)
    if engine not in {"auto", "tesseract"}:
        return "", None, "OCR_ENGINE must be auto, tesseract, or paddle."
    tesseract_path = os.getenv("TESSERACT_CMD")
    tesseract_found = bool(tesseract_path and Path(tesseract_path).is_file()) or bool(shutil.which("tesseract"))
    if not tesseract_found:
        for candidate in (r"C:\Program Files\Tesseract-OCR\tesseract.exe",
                          r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"):
            if Path(candidate).is_file():
                tesseract_found = True
                break
    if engine == "auto" and not tesseract_found:
        return _ocr_image_paddle(image)
    text, metrics, error = _ocr_image_tesseract(image)
    if metrics:
        metrics["engine"] = "Tesseract"
    return text, metrics, error


def _rule_line_counts(data: bytes):
    """Count strong straight ruling lines as a layout observation only."""
    try:
        from PIL import Image, ImageOps
        image = Image.open(io.BytesIO(data)).convert("L")
        image.thumbnail((1000, 1000))
        width, height = image.size
        pixels = list(image.getdata())
        def count_runs(values):
            marked = [sum(v < 110 for v in row) / len(row) >= .68 for row in values]
            runs, active = 0, False
            for yes in marked:
                if yes and not active: runs += 1
                active = yes
            return runs
        horizontal = count_runs([pixels[y*width:(y+1)*width] for y in range(height)])
        vertical = count_runs([[pixels[y*width+x] for y in range(height)] for x in range(width)])
        return horizontal, vertical
    except Exception:
        return None


def _border_geometry(data: bytes):
    """Measure long straight rules; this is a layout cue, not template verification."""
    try:
        import cv2
        import numpy as np
        image = cv2.imdecode(np.frombuffer(data, dtype=np.uint8), cv2.IMREAD_GRAYSCALE)
        if image is None:
            return {"available": False, "evidence": ["Image could not be decoded for border analysis."]}
        h, w = image.shape
        scale = min(1.0, 1800 / max(h, w))
        if scale < 1:
            image = cv2.resize(image, (int(w * scale), int(h * scale)), interpolation=cv2.INTER_AREA)
        h, w = image.shape
        edges = cv2.Canny(cv2.GaussianBlur(image, (3, 3), 0), 45, 130)
        min_length = max(30, int(min(w, h) * .18))
        lines = cv2.HoughLinesP(edges, 1, np.pi / 180, threshold=max(25, min_length // 3),
            minLineLength=min_length, maxLineGap=max(8, min_length // 10))
        horizontal, vertical = [], []
        if lines is not None:
            for line in lines[:, 0, :]:
                x1, y1, x2, y2 = map(int, line)
                dx, dy = x2 - x1, y2 - y1
                length = float((dx * dx + dy * dy) ** .5)
                if length < min_length:
                    continue
                angle = abs(float(np.degrees(np.arctan2(dy, dx))))
                if angle <= 12 or angle >= 168:
                    horizontal.append((x1, y1, x2, y2, angle if angle <= 90 else 180-angle))
                elif 78 <= angle <= 102:
                    vertical.append((x1, y1, x2, y2, abs(90-angle)))
        h_angles = [item[4] for item in horizontal]
        v_angles = [item[4] for item in vertical]
        h_spread = max(h_angles) - min(h_angles) if len(h_angles) > 1 else 0.0
        v_spread = max(v_angles) - min(v_angles) if len(v_angles) > 1 else 0.0
        intersections = 0
        for hx1, hy1, hx2, hy2, _ in horizontal:
            hy = (hy1 + hy2) / 2
            for vx1, vy1, vx2, vy2, _ in vertical:
                vx = (vx1 + vx2) / 2
                if min(hx1, hx2)-8 <= vx <= max(hx1, hx2)+8 and min(vy1, vy2)-8 <= hy <= max(vy1, vy2)+8:
                    intersections += 1
        evidence = [f"Detected {len(horizontal)} long horizontal and {len(vertical)} long vertical line segments (OpenCV)."]
        evidence.append(f"Approximate horizontal/vertical line intersections: {intersections}.")
        if h_angles:
            evidence.append(f"Horizontal line angle variation: {h_spread:.1f} degrees; this can also result from camera perspective or a curved page.")
        if v_angles:
            evidence.append(f"Vertical line angle variation: {v_spread:.1f} degrees; this can also result from camera perspective or a curved page.")
        enough = len(horizontal) >= 2 and len(vertical) >= 2
        flagged = enough and (h_spread > 6 or v_spread > 6)
        if not enough:
            evidence.append("Not enough long lines were detected to assess border alignment reliably.")
        return {"available": True, "enough": enough, "flagged": flagged, "evidence": evidence,
                "horizontal_count": len(horizontal), "vertical_count": len(vertical),
                "intersection_count": intersections,
                "horizontal_spread": h_spread, "vertical_spread": v_spread}
    except ImportError:
        return {"available": False, "evidence": ["OpenCV is not installed; border geometry analysis is unavailable."]}
    except Exception as exc:
        return {"available": False, "evidence": [f"Border analysis could not complete ({type(exc).__name__})."]}


def _jpeg_ela(data: bytes):
    """Basic error-level screening. ELA outliers are only leads for human review."""
    try:
        from PIL import Image, ImageChops, ImageStat
        source = Image.open(io.BytesIO(data))
        source.draft("RGB", (1800, 1800))
        source = source.convert("RGB")
        source.thumbnail((1800, 1800), Image.Resampling.LANCZOS)
        encoded = io.BytesIO()
        source.save(encoded, format="JPEG", quality=90)
        encoded.seek(0)
        recompressed = Image.open(encoded).convert("RGB")
        difference = ImageChops.difference(source, recompressed).convert("L")
        tile = 48
        scores = []
        for top in range(0, source.height, tile):
            for left in range(0, source.width, tile):
                box = (left, top, min(left+tile, source.width), min(top+tile, source.height))
                score = ImageStat.Stat(difference.crop(box)).mean[0]
                scores.append((score, box))
        ordered = sorted(score for score, _ in scores)
        median = ordered[len(ordered)//2] if ordered else 0
        threshold = max(12.0, median * 2.5 + 4)
        regions = []
        if median > 1:
            for score, box in scores:
                if score >= threshold:
                    x0, y0, x1, y1 = box
                    regions.append({"x_percent": round(x0/source.width*100, 1), "y_percent": round(y0/source.height*100, 1),
                                    "width_percent": round((x1-x0)/source.width*100, 1), "height_percent": round((y1-y0)/source.height*100, 1),
                                    "residual": round(score, 1)})
        return {"available": True, "median_residual": round(median, 2), "regions": regions[:12]}
    except Exception as exc:
        return {"available": False, "median_residual": None, "regions": [], "error": type(exc).__name__}


def _money_units(rupees: str, paise: str) -> int:
    return int(rupees.replace(",", "")) * 100 + int(paise)


def _numeric_row_pair(row):
    """Recover rupees/paise even when OCR splits a handwritten number into tokens."""
    digits = "".join(re.findall(r"\d", "".join(row.get("tokens", []))))
    if len(digits) < 3:
        return None
    return digits[:-2], digits[-2:]


def _check_receipt_total(text: str, numeric_lines=None, image_height=None, infer_footer_total=False):
    """Compare OCR-readable row amounts with a printed receipt total.

    Only handles common Indian receipt layouts where rupees and paise are the
    final two number columns. Treats every mismatch as a review signal, never
    proof of falsification.
    """
    lines = [re.sub(r"\s+", " ", row).strip() for row in text.splitlines() if row.strip()]
    total_label_present = any(re.search(r"\btotal\b", line, re.I) for line in lines)
    total_value = None
    total_line = ""
    footer_total_y = None
    # Printed decimal amounts (e.g. 7,895.53) or split rupee/paise columns (7895 53).
    for line in lines:
        m = re.search(r"\b(?:grand\s+|net\s+)?total\b\s*[:₹Rs.]*\s*([\d,]+)(?:\s*[.,]\s*(\d{2})|\s+(\d{2}))\s*$", line, re.I)
        if m:
            total_value = _money_units(m.group(1), m.group(2) or m.group(3))
            total_line = line
            break
    if total_value is None:
        # Digit-focused OCR may read the footer value while prose OCR reads
        # the printed TOTAL label. Only infer the footer pair when both exist.
        footer_threshold = .72 if total_label_present else .80
        footer = [row for row in (numeric_lines or [])
                  if image_height and row["top"] >= image_height * footer_threshold and row.get("tokens")]
        if (total_label_present or infer_footer_total) and footer:
            pair = _numeric_row_pair(footer[-1])
            if pair:
                total_value = _money_units(pair[0], pair[1])
                footer_total_y = footer[-1]["top"]
                total_line = ("TOTAL " if total_label_present else "Possible footer total (label not OCR-read) ") + " ".join(pair)
        if total_value is None:
            return None
    if footer_total_y is None and (total_label_present or infer_footer_total):
        threshold = .72 if total_label_present else .80
        footer = [x for x in (numeric_lines or [])
                  if image_height and x["top"] >= image_height * threshold and x.get("tokens")]
        if footer:
            footer_total_y = footer[-1]["top"]

    rows = []
    # Receipts often number item rows and print amount as the last two columns.
    for line in lines:
        if re.match(r"^(?:grand\s+|net\s+)?total\b", line, re.I):
            continue
        if not re.match(r"^\d{1,2}\s+\S", line):
            continue
        m = re.search(r"([\d,]+)\s+(\d{2})\s*$", line)
        if m:
            rows.append(_money_units(m.group(1), m.group(2)))
    numeric_evidence = False
    if len(rows) < 3 and numeric_lines:
        numeric_evidence = True
        # The table's amount columns occupy the right side; ignore the lower
        # footer row when a printed total label is present.
        for row in numeric_lines:
            if footer_total_y is not None and row["top"] == footer_total_y:
                continue
            if image_height and row["top"] >= image_height * .90:
                continue
            pair = _numeric_row_pair(row)
            if pair:
                rows.append(_money_units(pair[0], pair[1]))
    if len(rows) < 3:
        return {"available": False, "message": "A printed total was found, but OCR did not reliably identify at least three item amounts.", "evidence": [total_line]}
    item_sum = sum(rows)
    difference = total_value - item_sum
    if abs(difference) <= 2:
        return {"available": True, "match": True, "message": "OCR-readable item amounts add up to the printed total. This arithmetic check does not verify the issuer or document authenticity.",
                "evidence": [f"Items: ₹{item_sum // 100:,}.{item_sum % 100:02d} across {len(rows)} rows.", f"Printed total: ₹{total_value // 100:,}.{total_value % 100:02d}."]}
    return {"available": True, "match": False, "message": "OCR-readable item amounts do not add up to the printed total. Recheck the figures manually; OCR mistakes, discounts, taxes, or omitted lines may explain the difference.",
            "evidence": [f"Items: ₹{item_sum // 100:,}.{item_sum % 100:02d} across {len(rows)} rows.", f"Printed total: ₹{total_value // 100:,}.{total_value % 100:02d}.", f"Difference (printed total minus item sum): ₹{difference // 100:,}.{abs(difference) % 100:02d}" if difference >= 0 else f"Difference (printed total minus item sum): -₹{abs(difference) // 100:,}.{abs(difference) % 100:02d}"]}


def analyze(filename: str, data: bytes):
    ext = Path(filename).suffix.lower()
    findings, text, metadata, fonts = [], "", {}, []
    metadata_analysis = {}
    ocr_errors = []
    ocr_metrics = []
    numeric_lines = []
    ocr_image_height = None
    layout_spans = []
    image_layout = None
    border_analysis = None
    tamper_analysis = None
    forensics_image_data = None
    tamper_model_results = {}
    pdf_rule_count = 0
    doc_type = {".pdf": "PDF document", ".png": "PNG image", ".jpg": "JPEG image",
                ".jpeg": "JPEG image", ".docx": "Word document (DOCX)", ".doc": "Legacy Word document (DOC)"}.get(ext, "Unknown")

    if ext == ".pdf":
        try:
            import fitz
            pdf = fitz.open(stream=data, filetype="pdf")
            metadata = {k: str(v) for k, v in (pdf.metadata or {}).items() if v}
            metadata_analysis = _metadata_analysis(data, "PDF document", metadata)
            metadata_analysis["metadata_present"] = bool(metadata)
            metadata_analysis["metadata_keys"] = sorted(metadata.keys())
            pages = min(len(pdf), MAX_PAGES)
            for page_index, page in enumerate(pdf[:pages]):
                if page_index == 0:
                    forensics_image_data = page.get_pixmap(matrix=fitz.Matrix(1.25, 1.25), alpha=False).tobytes("png")
                text += page.get_text() + "\n"
                try:
                    pdf_rule_count += sum(1 for drawing in page.get_drawings() if drawing.get("items"))
                except Exception:
                    pass
                for block in page.get_text("dict").get("blocks", []):
                    for line in block.get("lines", []):
                        for span in line.get("spans", []):
                            if span.get("text", "").strip():
                                layout_spans.append({"text": span["text"].strip(), "size": float(span.get("size", 0)),
                                                     "font": str(span.get("font", ""))})
                for font in page.get_fonts(full=True):
                    fonts.append(str(font[3] or font[0]))
                if not page.get_text().strip():
                    pix = page.get_pixmap(matrix=fitz.Matrix(1.6, 1.6), alpha=False)
                    ocr, metrics, ocr_error = _ocr_image(pix.tobytes("png"))
                    text += ocr + "\n"
                    if ocr_error: ocr_errors.append(ocr_error)
                    if metrics: ocr_metrics.append(metrics)
                    if metrics:
                        numeric_lines.extend(metrics.get("numeric_lines", []))
                        ocr_image_height = metrics.get("height")
            findings.append(finding("PDF text and pages", "Passed" if text.strip() else "Warning",
                f"Inspected {pages} of {len(pdf)} pages." + (" Page limit reached." if len(pdf) > pages else ""),
                [f"Extracted {len(text.strip())} characters."]))
            findings.append(finding("Metadata analysis", "Not Checked",
                "PDF metadata fields were inventoried. Values can remain, change, be copied or be removed; they do not prove origin.",
                [f"{label}: {value}" for label, value in metadata_analysis.items() if label not in {"metadata_keys", "metadata_present"}], "low"))
            findings.append(finding("Embedded fonts", "Warning" if len(set(fonts)) > 1 else ("Passed" if fonts else "Not Checked"),
                "Multiple font resources occur in the PDF. This can be normal and is not evidence of alteration." if len(set(fonts)) > 1 else "No unusual font variation detected by this basic check." if fonts else "No embedded font information available.",
                sorted(set(fonts))[:10], "low"))
        except Exception as exc:
            findings.append(finding("PDF parsing", "Failed", "The PDF could not be safely parsed.", [type(exc).__name__], "high"))
    elif ext in {".png", ".jpg", ".jpeg"}:
        try:
            from PIL import Image
            image = Image.open(io.BytesIO(data))
            image.verify()
            image = Image.open(io.BytesIO(data))
            analysis_image_data = _analysis_raster_bytes(data, ext)
            forensics_image_data = analysis_image_data
            metadata = {str(k): str(v)[:300] for k, v in image.info.items() if k not in {"icc_profile", "exif"}}
            try:
                from PIL.ExifTags import TAGS
                for tag_id, value in image.getexif().items():
                    tag = TAGS.get(tag_id, str(tag_id))
                    if tag in {"Make", "Model", "Software", "DateTime", "DateTimeOriginal", "Orientation"}:
                        metadata[f"EXIF {tag}"] = str(value)[:200]
                    elif tag == "GPSInfo":
                        metadata["EXIF GPS"] = "Location metadata present; values omitted from the report."
            except Exception:
                pass
            text, metrics, ocr_error = _ocr_image(analysis_image_data)
            if ocr_error: ocr_errors.append(ocr_error)
            if metrics: ocr_metrics.append(metrics)
            if metrics:
                numeric_lines.extend(metrics.get("numeric_lines", []))
                ocr_image_height = metrics.get("height")
            image_layout = _rule_line_counts(analysis_image_data)
            border_analysis = _border_geometry(analysis_image_data)
            tamper_analysis = _jpeg_ela(data) if ext in {".jpg", ".jpeg"} else {"available": False, "regions": []}
            metadata_analysis = _image_metadata(data, image, ext)
            metadata = {"embedded_metadata_keys": ", ".join(metadata_analysis["metadata_keys"])} if metadata_analysis["metadata_present"] else {}
            ocr_status = "Passed" if text.strip() else "Not Checked" if ocr_error else "Warning"
            findings.append(finding("Image OCR", ocr_status,
                "Text was extracted with local OCR; this does not verify authenticity." if text.strip() else ocr_error or "OCR ran but returned no readable text.",
                [f"Extracted {len(text.strip())} characters."] + ([f"OCR engine: {metrics.get('engine', 'Tesseract')}; confidence is readability only, not an authenticity score." ] if metrics else []) + ([ocr_error] if ocr_error else [])))
            findings.append(finding("Metadata analysis", "Not Checked",
                "Embedded metadata fields were inventoried. Metadata can be preserved, altered, stripped or copied; it cannot prove authenticity.",
                [f"{label}: {value}" for label, value in metadata_analysis.items() if label not in {"metadata_keys", "metadata_present"}], "low"))
            claimed_software = metadata_analysis.get("software_used", "")
            editor_names = ("photoshop", "gimp", "photopea", "paint.net", "affinity photo", "canva")
            if claimed_software != "Not present in embedded metadata" and any(
                    editor in claimed_software.lower() for editor in editor_names):
                findings.append(finding("Image metadata editing-software clue", "Warning",
                    "Embedded metadata names image-editing software. This is a review clue only: it does not prove the document content was changed, and metadata can be copied or edited.",
                    [f"Software tag: {claimed_software}", "Request an issuer copy or compare against a trusted original."], "medium"))
        except Exception:
            findings.append(finding("Image parsing", "Failed", "The image is corrupt or unsupported.", severity="high"))
    elif ext == ".doc":
        findings.append(finding("Legacy DOC parsing", "Not Checked", "This legacy binary format needs an isolated conversion service, which is not configured."))
    elif ext == ".docx":
        try:
            import zipfile
            with zipfile.ZipFile(io.BytesIO(data)) as package:
                infos = package.infolist()
                if len(infos) > 2000 or sum(x.file_size for x in infos) > 50 * 1024 * 1024:
                    raise ValueError("DOCX package exceeds safe expansion limits")
                if any(x.filename.startswith(("/", "\\")) or ".." in Path(x.filename).parts for x in infos):
                    raise ValueError("DOCX package contains an unsafe path")
                if "[Content_Types].xml" not in package.namelist() or "word/document.xml" not in package.namelist():
                    raise ValueError("Missing DOCX package parts")
            from docx import Document
            doc = Document(io.BytesIO(data))
            metadata = {"author": doc.core_properties.author, "created": str(doc.core_properties.created or ""),
                        "modified": str(doc.core_properties.modified or ""), "last_modified_by": doc.core_properties.last_modified_by}
            metadata_analysis = _metadata_analysis(data, "DOCX document", metadata)
            metadata_analysis["metadata_present"] = any(metadata.values())
            metadata_analysis["metadata_keys"] = sorted(k for k, v in metadata.items() if v)
            text = "\n".join(p.text for p in doc.paragraphs)
            for table in doc.tables:
                text += "\n" + "\n".join(" | ".join(c.text for c in row.cells) for row in table.rows)
            findings.append(finding("DOCX text", "Passed" if text.strip() else "Warning", "Parsed document paragraphs and tables.", [f"Extracted {len(text.strip())} characters."]))
            findings.append(finding("Metadata analysis", "Not Checked", "DOCX core properties were inventoried; these fields are self-reported and may be edited or removed.",
                [f"{label}: {value}" for label, value in metadata_analysis.items() if label not in {"metadata_keys", "metadata_present"}], "low"))
        except Exception as exc:
            findings.append(finding("DOCX parsing", "Failed", "The file is invalid or exceeds safe DOCX parsing limits.", [type(exc).__name__], "high"))

    clean = re.sub(r"\s+", " ", text).strip()
    lowered = clean.lower()
    if ext in {".png", ".jpg", ".jpeg", ".pdf", ".docx"} and clean:
        pharmacy_header = any(k in lowered for k in ("pharmacy", "chemist"))
        medicine_content = bool(re.search(r"\b\d+(?:\.\d+)?\s?(?:mg|mcg|ml|tablet|tab)\b", lowered))
        if pharmacy_header and (medicine_content or any(k in lowered for k in ("qty", "quantity", "batch", "exp.date", "total"))):
            doc_type = "Pharmacy receipt / invoice (content-based estimate)"
        elif any(k in lowered for k in ("invoice", "bill to", "subtotal", "grand total")):
            doc_type = "Invoice / bill (content-based estimate)"
        elif any(k in lowered for k in ("certificate", "certifies that", "certification")):
            doc_type = "Certificate (content-based estimate)"
    layout_evidence = []
    if ocr_metrics:
        words = [word for metrics in ocr_metrics for word in metrics["words"]]
        line_count = sum(metrics["line_count"] for metrics in ocr_metrics)
        confs = [m["average_confidence"] for m in ocr_metrics if m["average_confidence"] is not None]
        if words:
            upper_limit = min(m["height"] for m in ocr_metrics) * .35
            top_words = [w for w in words if w["top"] < upper_limit]
            heading_word = max(top_words or words, key=lambda w: w["height"])
            heading = " ".join(w["text"] for w in sorted([w for w in (top_words or words)
                if abs(w["top"] - heading_word["top"]) <= heading_word["height"]], key=lambda w: w["left"]))
            layout_evidence.append(f"Probable prominent text near the top: {heading[:100]} (OCR-based estimate).")
        layout_evidence.append(f"OCR found approximately {line_count} text lines" + (f"; average word confidence {sum(confs)/len(confs):.0f}% (not a probability of authenticity)." if confs else "."))
    if layout_spans:
        biggest = max(layout_spans, key=lambda span: span["size"])
        layout_evidence.append(f"PDF text spans use {len({span['font'] for span in layout_spans})} font families; largest text span: {biggest['text'][:70]!r} ({biggest['size']:.1f} pt).")
    if pdf_rule_count:
        layout_evidence.append(f"PDF contains {pdf_rule_count} vector drawing groups that may include rules, boxes or borders; this check does not identify whether they match the issuer's template.")
    if image_layout is not None:
        layout_evidence.append(f"Image scan detected about {image_layout[0]} strong horizontal and {image_layout[1]} vertical ruling lines.")
    if border_analysis and border_analysis.get("available"):
        layout_evidence.extend(border_analysis.get("evidence", []))
        border_status = "Warning" if border_analysis.get("flagged") else "Passed" if border_analysis.get("enough") else "Not Checked"
        findings.append(finding("Border line alignment", border_status,
            "Long rule lines were measured for local straightness and angle spread. This is not a comparison with the issuer's original template and cannot establish authenticity.",
            border_analysis.get("evidence", []), "low" if border_analysis.get("flagged") else "info"))
    elif ext in {".jpg", ".jpeg", ".png"}:
        findings.append(finding("Border line alignment", "Not Checked",
            "Border geometry could not be measured. Install the OpenCV dependency and try again." if not border_analysis else "Border geometry could not be measured.",
            (border_analysis or {}).get("evidence", [])))
    if layout_evidence:
        readable_layout = bool(ocr_metrics or layout_spans or pdf_rule_count)
        message = ("Extracted structural cues for review. They were not compared with an issuer template; design differences alone do not indicate tampering." if readable_layout
                   else "Only basic border-line cues were inspected; OCR is unavailable, so headings and text layout could not be checked.")
        # Structural cues are observations only. Without a trusted issuer
        # template this check did not verify a document's formatting.
        findings.append(finding("Formatting, heading and borders", "Not Checked", message, layout_evidence))
    elif ext in {".png", ".jpg", ".jpeg", ".pdf"}:
        findings.append(finding("Formatting, heading and borders", "Not Checked",
            "OCR or readable text layout was unavailable, so headings, spacing and borders could not be inspected."))
    if not findings or not any(x["status"] == "Failed" for x in findings):
        if not clean:
            findings.append(finding("Readable content", "Not Checked" if ocr_errors else "Warning",
                "Content checks were not run because OCR is unavailable. " + "; ".join(dict.fromkeys(ocr_errors)) if ocr_errors else "OCR ran but no readable text was found.", severity="low"))
        else:
            repeated = [s for s, n in Counter(re.findall(r"\b[\w@./-]{4,}\b", clean.lower())).items() if n >= 5]
            findings.append(finding("Text consistency", "Warning" if repeated else "Not Checked",
                "Repeated tokens detected by a basic heuristic; repetition can be legitimate and needs review." if repeated else "No repeated-token issue was detected by this limited check; this does not validate the document's text.", repeated[:8], "low"))
            dates = re.findall(r"\b(?:\d{1,2}[/-]){2}\d{2,4}\b|\b\d{4}-\d{2}-\d{2}\b", clean)
            bad_dates = []
            for value in dates:
                for fmt in ("%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y", "%m/%d/%y"):
                    try:
                        datetime.strptime(value, fmt)
                        break
                    except ValueError:
                        continue
                else:
                    bad_dates.append(value)
            findings.append(finding("Date format check", "Warning" if bad_dates else "Not Checked",
                "Some date-shaped values are impossible or unrecognized." if bad_dates else "Extracted date-shaped values appear calendar-valid; expiry, chronology, and issuer validity were not checked.", bad_dates[:8], "medium" if bad_dates else "info"))
            receipt_like = doc_type.startswith(("Pharmacy receipt", "Invoice / bill"))
            if re.search(r"\btotal\b", text, re.I) or receipt_like:
                arithmetic = _check_receipt_total(text, numeric_lines, ocr_image_height,
                    infer_footer_total=doc_type.startswith("Pharmacy receipt"))
                if arithmetic is None:
                    findings.append(finding("Receipt arithmetic", "Not Checked", "Could not read a printed total and enough item amounts to compare. Handwritten numbers were missed or ambiguous in OCR; check the figures manually. This check cannot determine authenticity."))
                elif not arithmetic["available"]:
                    findings.append(finding("Receipt arithmetic", "Not Checked", arithmetic["message"], arithmetic["evidence"]))
                else:
                    findings.append(finding("Receipt arithmetic", "Warning" if not arithmetic["match"] else "Passed",
                        arithmetic["message"], arithmetic["evidence"], "medium" if not arithmetic["match"] else "info"))
    consistency_checks = [x for x in findings if x["name"] in {"Text consistency", "Date format check", "Receipt arithmetic"}]
    consistency_warning = any(x["status"] == "Warning" and x["severity"] in {"medium", "high"} for x in consistency_checks)
    consistency_passed = any(x["status"] == "Passed" for x in consistency_checks)
    if consistency_warning:
        consistency_status = "Warning"
        consistency_message = "At least one cross-field or arithmetic check needs manual review. This does not confirm forgery."
        consistency_severity = "medium"
    elif consistency_passed:
        consistency_status = "Passed"
        consistency_message = "Some limited consistency checks completed without finding a mismatch; this does not verify the issuer or authenticity."
        consistency_severity = "info"
    else:
        consistency_status = "Not Checked"
        consistency_message = "There was not enough readable information to run document consistency checks."
        consistency_severity = "info"
    findings.append(finding("Document consistency", consistency_status, consistency_message,
        [f"{x['name']}: {x['message']}" for x in consistency_checks], consistency_severity))
    def labeled_values(labels):
        values = re.findall(rf"(?im)^\s*(?:{labels})\s*:?\s*([^\r\n]+)", text)
        return [re.sub(r"[.·_]{3,}.*$", "", v).strip(" .:·")[:160] for v in values if v.strip(" .:·")][:10]
    heading_lines = [line.strip() for line in text.splitlines() if len(line.strip()) >= 5 and sum(ch.isalpha() for ch in line) >= 5 and line.strip().upper() == line.strip()]
    amount_values = list(dict.fromkeys(re.findall(r"(?:[$€£₹]\s?\d[\d,]*(?:\.\d{1,2})?|\b\d[\d,]*\.\d{2}\b)", clean)))
    for line in text.splitlines():
        if re.match(r"^\s*\d{1,2}\s+\S", line):
            pair = re.search(r"([\d,]+)\s+(\d{2})\s*$", line)
            if pair:
                amount_values.append(f"₹{int(pair.group(1).replace(',', '')):,}.{pair.group(2)}")
        total_pair = re.search(r"\b(?:grand\s+|net\s+)?total\b\s*[:₹Rs.]*\s*([\d,]+)(?:\s*[.,]\s*(\d{2})|\s+(\d{2}))\s*$", line, re.I)
        if total_pair:
            paise = total_pair.group(2) or total_pair.group(3)
            amount_values.append(f"₹{int(total_pair.group(1).replace(',', '')):,}.{paise}")
    for row in numeric_lines:
        pair = _numeric_row_pair(row)
        if pair:
            try:
                cents = _money_units(pair[0], pair[1])
                amount_values.append(f"₹{cents // 100:,}.{cents % 100:02d}")
            except ValueError:
                pass
    fields = {"dates": list(dict.fromkeys(re.findall(r"\b(?:\d{1,2}[/-]){2}\d{2,4}\b|\b\d{4}-\d{2}-\d{2}\b", clean)))[:20],
              "amounts": list(dict.fromkeys(amount_values))[:20],
              "emails": list(dict.fromkeys(re.findall(r"\b[\w.+-]+@[\w.-]+\.\w{2,}\b", clean)))[:20],
              "identifiers": list(dict.fromkeys(re.findall(r"\b[A-Z0-9][A-Z0-9/-]{6,}\b", text)))[:20],
              "headings": list(dict.fromkeys(heading_lines))[:5],
              "patient_names": labeled_values(r"patient\s+name|patient"),
              "prescriber_names": labeled_values(r"prescriber\s+name|doctor\s+name|physician"),
              "addresses": labeled_values(r"address")}
    text = text[:30000]
    if ext in {".jpg", ".jpeg"} and tamper_analysis and tamper_analysis.get("available"):
        regions = tamper_analysis.get("regions", [])
        evidence = [f"JPEG recompression screening average residual: {tamper_analysis['median_residual']:.2f} (not an authenticity score)."]
        evidence.extend(f"Candidate region near x={r['x_percent']}%, y={r['y_percent']}% of page ({r['width_percent']}% × {r['height_percent']}%); edges, text and camera noise can trigger this." for r in regions[:6])
        findings.append(finding("Image manipulation screening (ELA)", "Warning" if regions else "Not Checked",
            "Localized compression differences were observed; this screening can be triggered by normal text, edges, lighting or JPEG compression and does not prove editing." if regions else "No strong localized compression outlier was detected. This does not rule out editing.",
            evidence, "low"))
    elif ext in {".jpg", ".jpeg"}:
        findings.append(finding("Image manipulation screening (ELA)", "Not Checked", "JPEG compression screening could not be completed."))
    if forensics_image_data and ext in {".png", ".jpg", ".jpeg", ".pdf"}:
        from app.vision_plugins import run_local_screening_model
        for model_name, title in (("doctamper", "DocTamper document-text screening"), ("trufor", "TruFor image-forgery screening")):
            result = run_local_screening_model(model_name, forensics_image_data, ".png" if ext == ".pdf" else ext)
            tamper_model_results[model_name] = result
            score_note = f"Raw screening score: {result['screening_score']:.3f}; not a calibrated authenticity probability." if result.get("screening_score") is not None else "No calibrated authenticity probability is available."
            page_note = "First PDF page only. " if ext == ".pdf" else ""
            evidence = [page_note + score_note]
            evidence.extend(f"Candidate region x={r['x_percent']:.1f}%, y={r['y_percent']:.1f}% ({r['width_percent']:.1f}% × {r['height_percent']:.1f}%)." for r in result.get("regions", []))
            findings.append(finding(title, result["status"], page_note + result["message"], evidence,
                "medium" if result.get("anomaly_detected") else "info"))
    if not metadata_analysis:
        metadata_analysis = _metadata_analysis(data, doc_type, metadata)
        metadata_analysis["metadata_present"] = bool(metadata)
        metadata_analysis["metadata_keys"] = sorted(metadata.keys())
    findings += [finding("QR and barcode", "Not Checked", "QR/barcode decoding is not included in this phase."),
                 finding("Digital signature", "Not Checked", "Cryptographic signature integrity and trust chain were not checked."),
                 finding("Issuer verification", "Not Checked", "No official issuer service is connected; no external lookup was attempted."),
                 finding("Image manipulation", "Not Checked", "No pixel-level alteration can be confirmed by this prototype. Visual anomalies and metadata are only screening cues.")]
    # Parsing/OCR errors do not establish that an authenticity verification failed.
    serious = any(x["status"] == "Failed" and x["name"] in {"Digital signature", "Issuer verification"} for x in findings)
    suspicious = any(x["status"] == "Warning" and x["severity"] in {"medium", "high"} for x in findings)
    status = "Verification Failed" if serious else "Suspicious — Review Required" if suspicious else "Unable to Verify"
    return {"document_type": doc_type, "status": status, "confidence": None, "findings": findings,
            "fields": fields, "extracted_text": text, "metadata": metadata,
            "metadata_analysis": metadata_analysis,
            "image_forensics": {"border_geometry": border_analysis,
                                "ela_regions": (tamper_analysis or {}).get("regions", []),
                                "doctamper": tamper_model_results.get("doctamper"),
                                "trufor": tamper_model_results.get("trufor")},
            "limitations": ["Appearance and basic heuristics cannot prove authenticity.", "No trusted issuer or cryptographic verification source is configured.", "OCR errors and metadata manipulation are possible."]}
